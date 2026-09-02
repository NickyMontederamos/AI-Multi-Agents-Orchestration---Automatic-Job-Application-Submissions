#!/usr/bin/env python3
"""
Convert a career-ai resume or cover-letter Markdown draft into a submission-ready
.docx file. Markdown stays the editable source of truth (it's what
resume-ats-scorer-agent scores) — this script is the final export step, run once a
draft has cleared or capped out its revision cycles.

Supported Markdown subset (matches what career-resume-agent actually produces):
  # Name                    -> title line (bold, centered, larger)
  plain line right after H1 -> contact line (centered, smaller)
  ## Section Header         -> Heading 2 (standard ATS-safe section header)
  - bullet text             -> bulleted list item
  **bold**, *italic*        -> inline runs (non-nested)
  blank line                -> paragraph break
  everything else           -> normal left-aligned paragraph

Format-compliance choices are deliberate, matching resume-ats-scorer-agent's
checklist: single column, no tables/text boxes/headers-footers/graphics, one
standard font (Calibri), consistent styling throughout.

Usage:
  python md_to_docx.py <input.md> <output.docx>
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

INLINE_TOKEN = re.compile(r"(\*\*.+?\*\*|\*.+?\*)")


def add_inline_runs(paragraph, text, base_italic=False):
    for token in INLINE_TOKEN.split(text):
        if not token:
            continue
        if token.startswith("**") and token.endswith("**"):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
            run.italic = base_italic
        elif token.startswith("*") and token.endswith("*"):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        else:
            run = paragraph.add_run(token)
            run.italic = base_italic


def convert(md_path: Path, docx_path: Path) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()

    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(11)
    for section in doc.sections:
        section.left_margin = section.right_margin = Pt(54)  # ~0.75in

    i = 0
    n = len(lines)
    first_heading_seen = False

    while i < n:
        line = lines[i].rstrip()

        if not line:
            i += 1
            continue

        if line.startswith("# "):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(line[2:].strip())
            run.bold = True
            run.font.size = Pt(18)
            first_heading_seen = True
            i += 1
            # contact line immediately follows a resume's H1, if present
            if i < n and lines[i].strip() and not lines[i].startswith(("#", "-")):
                contact = doc.add_paragraph()
                contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = contact.add_run(lines[i].strip())
                run.font.size = Pt(10)
                i += 1
            continue

        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
            i += 1
            continue

        if line.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline_runs(p, line[2:].strip())
            i += 1
            continue

        # plain paragraph (cover letters, resume summary, italic role/date lines)
        p = doc.add_paragraph()
        add_inline_runs(p, line)
        i += 1

    docx_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(docx_path))


def main():
    if len(sys.argv) != 3:
        print("Usage: python md_to_docx.py <input.md> <output.docx>", file=sys.stderr)
        sys.exit(1)
    src, dst = Path(sys.argv[1]), Path(sys.argv[2])
    if not src.exists():
        print(f"Input not found: {src}", file=sys.stderr)
        sys.exit(1)
    convert(src, dst)
    print(f"Wrote {dst}")


if __name__ == "__main__":
    main()
